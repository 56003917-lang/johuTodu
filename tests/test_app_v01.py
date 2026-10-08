# -*- coding: utf-8 -*-
"""《安全生产在线培训平台 v0.1》程序自动化测试。

依据《技术设计文档》§12：学习引擎的计时口径（间隔边界、弹窗暂停、单日上限、
断网重放）全部写成自动化测试用例先行——这是全系统唯一"错不起"的模块。

运行：.venv\\Scripts\\python.exe tests\\test_app_v01.py
覆盖：
  配置/建库(WAL)/引导管理员/名册导入/登录锁定/首登改密门禁/会话Cookie
  课程扫描全生命周期/服务端权威计时/达标判定/断点续学
  防挂机状态机（含验收标准3：不响应弹窗计时暂停）
  单日上限/断线超长/时钟回拨/心跳超时清理/双开防护
  台账导出（xlsx+csv）/导出留痕/总览/管理员越权/Range 下载/备份
"""
import io
import json
import os
import shutil
import sqlite3
import sys
import tempfile
from datetime import timedelta
from pathlib import Path

REPO_ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(REPO_ROOT))

from fastapi.testclient import TestClient  # noqa: E402
from openpyxl import Workbook, load_workbook  # noqa: E402

from app.config import Settings  # noqa: E402
from app.main import create_app  # noqa: E402
from app.util import utcnow  # noqa: E402

ADMIN_NO, ADMIN_PW = "admin", "admin@2026"
LEARNER = {"employee_no": "A001", "name": "王建国", "department": "一项目部·高处作业班组", "role": "学员"}


# ====================== 测试基建 ======================
class FakeClock:
    """可注入的确定性时钟（测试计时口径专用）。"""

    def __init__(self):
        self.now = utcnow()

    def __call__(self):
        return self.now

    def advance(self, seconds):
        self.now = self.now + timedelta(seconds=seconds)


def make_pdf(pages=2, text="Safety Training Sample") -> bytes:
    """构造最小合法 PDF（无需外部库）。"""
    objs = []
    page_ids = []
    n_font = 3
    for i in range(pages):
        pid = len(objs) + 1
        cid = pid + 1
        page_ids.append(pid)
        objs.append(("<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] "
                     "/Resources << /Font << /F1 %d 0 R >> >> /Contents %d 0 R >>" % (n_font, cid)).encode())
        content = "BT /F1 16 Tf 72 700 Td (%s - Page %d) Tj ET" % (text, i + 1)
        objs.append(("<< /Length %d >>\nstream\n%s\nendstream" % (len(content), content)).encode())
    kids = " ".join("%d 0 R" % p for p in page_ids)
    objs.insert(0, b"<< /Type /Catalog /Pages 2 0 R >>")
    objs.insert(1, ("<< /Type /Pages /Kids [%s] /Count %d >>" % (kids, pages)).encode())
    objs.append(b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>")

    out = io.BytesIO()
    out.write(b"%PDF-1.4\n")
    offsets = []
    for i, body in enumerate(objs, start=1):
        offsets.append(out.tell())
        out.write(("%d 0 obj\n" % i).encode("ascii"))
        out.write(body)
        out.write(b"\nendobj\n")
    xref_pos = out.tell()
    out.write(("xref\n0 %d\n" % (len(objs) + 1)).encode("ascii"))
    out.write(b"0000000000 65535 f \n")
    for off in offsets:
        out.write(("%010d 00000 n \n" % off).encode("ascii"))
    out.write(("trailer << /Size %d /Root 1 0 R >>\nstartxref\n%d\n%%%%EOF" % (len(objs) + 1, xref_pos)).encode("ascii"))
    return out.getvalue()


def make_roster_xlsx(rows):
    wb = Workbook()
    ws = wb.active
    ws.append(["工号", "姓名", "部门/班组", "角色"])
    for r in rows:
        ws.append(r)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


class Env:
    """一个测试环境：临时数据目录 + 应用 + TestClient + 假时钟。

    :param roster_force_pwd: 是否启用首登强制改密；默认 False（测试直接学习，
        改密门禁由 test_password_change_gate / test_bootstrap_admin 专门覆盖）。
    """

    def __init__(self, study=None, registry=None, roster_force_pwd=False, default_minutes=None):
        self.tmp = Path(tempfile.mkdtemp(prefix="tp-test-", dir=str(REPO_ROOT / ".temp")))
        cfg = {"data_dir": str(self.tmp)}
        if study:
            cfg["study"] = study
        if roster_force_pwd is not None:
            cfg["roster"] = {"force_change_password": roster_force_pwd}
        courses_cfg = {}
        if registry is not None:
            courses_cfg["registry"] = registry
        if default_minutes is not None:
            courses_cfg["default_required_minutes"] = default_minutes
        if courses_cfg:
            cfg["courses"] = courses_cfg
        self.settings = Settings.from_dict(cfg, base_dir=str(self.tmp))
        self.clock = FakeClock()
        self.app = create_app(self.settings, clock=self.clock, start_tasks=False)
        # 以上下文方式启动，lifespan（建表/引导管理员）才会执行
        self.client = TestClient(self.app)
        self.client.__enter__()

    # ---- 便捷方法 ----
    def login(self, no, pw):
        return self.client.post("/api/v1/auth/login",
                                json={"employee_no": no, "password": pw}).json()

    def login_ok(self, no, pw):
        j = self.login(no, pw)
        assert j["code"] == 0, "登录失败: %s" % j
        return j["data"]

    def import_roster(self, rows):
        data = make_roster_xlsx(rows)
        r = self.client.post("/api/v1/admin/roster/import",
                             files={"file": ("roster.xlsx", data,
                                             "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")})
        j = r.json()
        assert j["code"] == 0, "名册导入失败: %s" % j
        return j["data"]

    def put_course(self, rel, content=None):
        p = self.settings.courses_dir / rel
        p.parent.mkdir(parents=True, exist_ok=True)
        p.write_bytes(content if content is not None else make_pdf())
        return p

    def scan(self):
        r = self.client.post("/api/v1/admin/courses/scan")
        j = r.json()
        assert j["code"] == 0, "扫描失败: %s" % j
        return j["data"]

    def course_id(self, rel):
        # 直接查库：不依赖当前会话角色（学员/管理员阶段均可调用）
        import sqlite3 as _s
        conn = _s.connect(str(self.settings.db_path))
        try:
            row = conn.execute("SELECT id FROM courses WHERE file_name=?", (rel,)).fetchone()
        finally:
            conn.close()
        if not row:
            raise AssertionError("课程不存在: %s" % rel)
        return row[0]

    def start(self, course_id):
        r = self.client.post("/api/v1/study/start", json={"course_id": course_id})
        j = r.json()
        assert j["code"] == 0, "开课失败: %s" % j
        return j["data"]

    def hb(self, sid, page=1, visible=True):
        r = self.client.post("/api/v1/study/heartbeat",
                             json={"session_id": sid, "last_position": {"page": page, "ratio": None},
                                   "visible": visible})
        j = r.json()
        assert j["code"] == 0, "心跳失败: %s" % j
        return j["data"]

    def confirm(self, sid, token):
        return self.client.post("/api/v1/study/confirm",
                                json={"session_id": sid, "token": token}).json()

    def stop(self, sid, page=5):
        r = self.client.post("/api/v1/study/stop",
                             json={"session_id": sid, "last_position": {"page": page, "ratio": None},
                                   "reason": "normal"})
        j = r.json()
        assert j["code"] == 0, "结束失败: %s" % j
        return j["data"]

    def standard_learner(self, courses=None):
        """管理员阶段完成名册导入与课程上架，再切换到学员 A001 会话（免改密）。

        :param courses: 需预先放入课程库的课程，列表（默认内容）或 dict（rel→自定义内容）。
        """
        self.login_ok(ADMIN_NO, ADMIN_PW)
        self.import_roster([[LEARNER["employee_no"], LEARNER["name"], LEARNER["department"], ""]])
        items = list((courses or {}).items()) if isinstance(courses, dict) else [(rel, None) for rel in (courses or [])]
        for rel, content in items:
            self.put_course(rel, content)
        if items:
            self.scan()
        self.client.post("/api/v1/auth/logout")
        self.login_ok("A001", "A001")
        return "A001"

    def close(self):
        try:
            self.client.__exit__(None, None, None)
        finally:
            self.app.state.db.dispose()
            shutil.rmtree(self.tmp, ignore_errors=True)


def with_env(fn):
    def wrapper():
        env = Env()
        try:
            fn(env)
        finally:
            env.close()
    wrapper.__name__ = fn.__name__
    return wrapper


# ====================== 基础设施 ======================
@with_env
def test_healthz_and_pages(env):
    r = env.client.get("/healthz")
    assert r.status_code == 200 and r.json()["status"] == "ok" and r.json()["db"] is True
    for path, kw in [("/", "员工登录"), ("/admin", "管理员登录")]:
        r = env.client.get(path)
        assert r.status_code == 200 and kw in r.text, "页面 %s 内容异常" % path
        assert "http://" not in r.text.replace("http://127.0.0.1", ""), "页面不应引用外部资源"


def test_config_defaults_per_doc():
    s = Settings.from_dict({})
    st = s.study
    assert (st.heartbeat_interval_s, st.max_gap_s, st.challenge_min_s,
            st.challenge_max_s, st.challenge_timeout_s, st.daily_cap_hours) == (15, 60, 90, 180, 120, 8)
    assert s.allowed_ext == ("pdf",) and s.max_file_mb == 100 and s.scan_interval_min == 5
    assert s.initial_password == "employee_no" and s.force_change_password is True
    assert s.backup_daily_at == "02:30" and s.backup_retain_days == 30


def test_config_yaml_load():
    p = REPO_ROOT / ".temp" / "cfg-test.yaml"
    p.parent.mkdir(exist_ok=True)
    p.write_text("study:\n  heartbeat_interval_s: 5\n  daily_cap_hours: 2\n"
                 "courses:\n  registry:\n    a.pdf:\n      required_minutes: 30\n", encoding="utf-8")
    s = Settings.load(p)
    assert s.study.heartbeat_interval_s == 5 and s.study.daily_cap_s == 7200
    assert s.registry_lookup("a.pdf", "a")["required_minutes"] == 30
    p.unlink()


@with_env
def test_db_init_wal_and_tables(env):
    conn = sqlite3.connect(str(env.settings.db_path))
    tables = {r[0] for r in conn.execute("SELECT name FROM sqlite_master WHERE type='table'")}
    conn.close()
    for t in ("users", "courses", "study_records", "study_sessions", "export_logs"):
        assert t in tables, "缺少表 %s" % t
    conn = sqlite3.connect(str(env.settings.db_path))
    mode = conn.execute("PRAGMA journal_mode").fetchone()[0]
    conn.close()
    assert mode.lower() == "wal", "WAL 模式未生效: %s" % mode


def test_bootstrap_admin():
    # 默认配置：引导管理员首登强制改密（技术文档 §9：初始密码=工号必须改）
    e1 = Env(roster_force_pwd=True)
    try:
        j = e1.login(ADMIN_NO, ADMIN_PW)
        assert j["code"] == 0 and j["data"]["role"] == "admin" and j["data"]["must_change_pwd"] is True, j
    finally:
        e1.close()
    # 免改密配置：must_change_pwd=False
    e2 = Env(roster_force_pwd=False)
    try:
        j = e2.login(ADMIN_NO, ADMIN_PW)
        assert j["code"] == 0 and j["data"]["must_change_pwd"] is False, j
    finally:
        e2.close()


@with_env
def test_auth_session_and_envelope(env):
    r = env.client.get("/api/v1/me")
    assert r.status_code == 401 and r.json()["code"] == 4101, "未登录应返回 4101"
    assert set(r.json().keys()) == {"code", "data", "msg"}, "响应未按统一包裹结构"
    j = env.login(ADMIN_NO, ADMIN_PW)
    assert j["code"] == 0 and j["msg"] == "" and j["data"]["employee_no"] == ADMIN_NO
    cookie = env.client.cookies.get("tp_session")
    assert cookie, "登录未下发会话 Cookie"
    # Cookie 安全属性（HttpOnly/SameSite）
    r = env.client.post("/api/v1/auth/login", json={"employee_no": ADMIN_NO, "password": ADMIN_PW})
    sc = r.headers.get("set-cookie", "")
    assert "httponly" in sc.lower() and "samesite=lax" in sc.lower(), "Cookie 安全属性缺失: %s" % sc
    me = env.client.get("/api/v1/me").json()
    assert me["code"] == 0 and me["data"]["role"] == "admin"
    # 登出后会话失效
    env.client.post("/api/v1/auth/logout")
    r = env.client.get("/api/v1/me")
    assert r.json()["code"] == 4101, "登出后仍可访问"


@with_env
def test_login_lockout(env):
    for i in range(5):
        j = env.login(ADMIN_NO, "wrong")
        assert j["code"] == 4301
    # 第 5 次失败后锁定：正确密码也被拒
    j = env.login(ADMIN_NO, ADMIN_PW)
    assert j["code"] == 4401, "连续 5 次失败未锁定: %s" % j
    # 锁定 10 分钟后恢复
    env.clock.advance(601)
    j = env.login(ADMIN_NO, ADMIN_PW)
    assert j["code"] == 0, "锁定期过后应可登录: %s" % j


# ====================== 名册 ======================
@with_env
def test_roster_template_and_import(env):
    env.login_ok(ADMIN_NO, ADMIN_PW)
    r = env.client.get("/api/v1/admin/roster/template")
    assert r.status_code == 200 and "spreadsheetml" in r.headers["content-type"]
    wb = load_workbook(io.BytesIO(r.content))
    ws = wb.active
    assert [c.value for c in ws[1][:4]] == ["工号", "姓名", "部门/班组", "角色"], "模板表头不符"

    stats = env.import_roster([
        ["A001", "王建国", "一项目部", ""],
        ["A002", "李安全", "安全管理中心", "管理员"],
        ["", "缺工号行", "x", ""],           # 无效行
    ])
    assert stats == {"total": 3, "created": 2, "updated": 0, "skipped": 1}, stats
    # 学员初始密码=工号
    j = env.login("A001", "A001")
    assert j["code"] == 0 and j["data"]["role"] == "learner"
    j = env.login("A002", "A002")
    assert j["code"] == 0 and j["data"]["role"] == "admin", "角色列应支持管理员"
    # 重复导入=更新，不改密码
    env.client.post("/api/v1/auth/logout")
    env.login_ok("A002", "A002")
    stats = env.import_roster([["A001", "王建国改", "二项目部", ""]])
    assert stats["created"] == 0 and stats["updated"] == 1, stats
    j = env.login("A001", "A001")
    assert j["code"] == 0 and j["data"]["name"] == "王建国改", "更新后姓名未生效且密码不应被重置"
    # 管理员接口学员不可用
    env.client.post("/api/v1/auth/logout")
    env.login_ok("A001", "A001")
    r = env.client.get("/api/v1/admin/overview")
    assert r.json()["code"] == 4501, "学员访问管理端应 4501"


@with_env
def test_password_change_gate(env):
    env3 = Env(roster_force_pwd=True)
    try:
        # 管理员阶段：导入名册 + 上架课程
        env3.login_ok(ADMIN_NO, ADMIN_PW)
        env3.import_roster([[LEARNER["employee_no"], LEARNER["name"], LEARNER["department"], ""]])
        env3.put_course("安全通识.pdf")
        env3.scan()
        cid = env3.course_id("安全通识.pdf")
        # 切到学员：首登未改密，学习被门禁拦截（4304）
        env3.client.post("/api/v1/auth/logout")
        env3.login_ok("A001", "A001")
        r = env3.client.post("/api/v1/study/start", json={"course_id": cid})
        assert r.json()["code"] == 4304, "首登未改密应拦截学习: %s" % r.json()
        # 改密：旧密码错误 → 4301
        r = env3.client.post("/api/v1/me/password", json={"old_password": "bad", "new_password": "abc123"})
        assert r.json()["code"] == 4301
        # 新密码过短 → 4001
        r = env3.client.post("/api/v1/me/password", json={"old_password": "A001", "new_password": "123"})
        assert r.json()["code"] == 4001
        # 改密成功后可学习
        r = env3.client.post("/api/v1/me/password", json={"old_password": "A001", "new_password": "a001@xyz"})
        assert r.json()["code"] == 0
        r = env3.client.post("/api/v1/study/start", json={"course_id": cid})
        assert r.json()["code"] == 0
        # 旧密码不再有效，新密码可登录
        env3.client.post("/api/v1/auth/logout")
        assert env3.login("A001", "A001")["code"] == 4301
        assert env3.login("A001", "a001@xyz")["code"] == 0
    finally:
        env3.close()


# ====================== 课程扫描 ======================
@with_env
def test_course_scan_lifecycle(env):
    env.login_ok(ADMIN_NO, ADMIN_PW)
    env.put_course("高处作业/登高作业安全要点.pdf")
    env.put_course("有限空间/人孔井作业安全规范.pdf")
    env.put_course("说明.txt", b"not pdf")
    stats = env.scan()
    assert stats["added"] == 2 and stats["skipped_ext"] == 1 and stats["active"] == 2, stats

    r = env.client.get("/api/v1/admin/courses").json()["data"]
    by_name = {c["file_name"]: c for c in r}
    assert by_name["高处作业/登高作业安全要点.pdf"]["category"] == "高处作业"
    assert by_name["高处作业/登高作业安全要点.pdf"]["required_minutes"] == 60  # 默认达标学时
    assert by_name["有限空间/人孔井作业安全规范.pdf"]["status"] == 1

    # registry 登记覆盖达标学时
    env2 = Env(registry={"人孔井作业安全规范.pdf": {"required_minutes": 30}}, default_minutes=45)
    try:
        env2.login_ok(ADMIN_NO, ADMIN_PW)
        env2.put_course("人孔井作业安全规范.pdf")
        env2.scan()
        assert env2.course_id("人孔井作业安全规范.pdf") > 0
        c = [x for x in env2.client.get("/api/v1/admin/courses").json()["data"]
             if x["file_name"] == "人孔井作业安全规范.pdf"][0]
        assert c["required_minutes"] == 30, "registry 未生效"
    finally:
        env2.close()

    # 文件替换：指纹更新、课程行保留（历史学时不丢）
    cid_before = env.course_id("高处作业/登高作业安全要点.pdf")
    env.put_course("高处作业/登高作业安全要点.pdf", make_pdf(pages=5, text="v2"))
    stats = env.scan()
    assert stats["updated"] == 1 and stats["added"] == 0, stats
    assert env.course_id("高处作业/登高作业安全要点.pdf") == cid_before, "替换文件不应换课程行"

    # 文件消失：下架但保留行
    (env.settings.courses_dir / "有限空间" / "人孔井作业安全规范.pdf").unlink()
    stats = env.scan()
    assert stats["removed"] == 1, stats
    c = [x for x in env.client.get("/api/v1/admin/courses").json()["data"]
         if x["file_name"] == "有限空间/人孔井作业安全规范.pdf"][0]
    assert c["status"] == 0, "文件消失后应下架"
    # 下架课程学员不可见、不可开课（名册仍在管理员会话内先导入）
    env.import_roster([[LEARNER["employee_no"], LEARNER["name"], LEARNER["department"], ""]])
    env.client.post("/api/v1/auth/logout")
    env.login_ok("A001", "A001")
    ids = [c["id"] for c in env.client.get("/api/v1/courses").json()["data"]]
    assert env.course_id("有限空间/人孔井作业安全规范.pdf") not in ids, "下架课程不应出现在学员列表"
    r = env.client.post("/api/v1/study/start", json={"course_id": env.course_id("有限空间/人孔井作业安全规范.pdf")})
    assert r.json()["code"] == 4041


# ====================== 学习引擎：计时口径 ======================
@with_env
def test_study_flow_and_completion(env):
    """验收标准 1：完整学习、学时自动累计、服务端权威。"""
    env.login_ok(ADMIN_NO, ADMIN_PW)
    env.import_roster([[LEARNER["employee_no"], LEARNER["name"], LEARNER["department"], ""]])
    env.put_course("通识/安全法解读.pdf")
    env.scan()
    cid = env.course_id("通识/安全法解读.pdf")
    env.client.post("/api/v1/auth/logout")
    env.login_ok("A001", "A001")

    data = env.start(cid)
    sid = data["session_id"]
    assert data["total_seconds"] == 0 and data["required_seconds"] == 3600 and data["last_position"] is None

    def hb_confirm(sid):
        """模拟认真学员：心跳一次，遇防挂机弹窗立即确认（真实页面行为）。"""
        d = env.hb(sid)
        if d.get("challenge"):
            j = env.confirm(sid, d["challenge"]["token"])
            assert j["code"] == 0, "弹窗确认失败: %s" % j
        return d

    # 正常心跳：15s 计 15s（服务端权威，累计值由服务端返回）
    env.clock.advance(15)
    d = hb_confirm(sid)
    assert d["gained_seconds"] == 15 and d["total_seconds"] == 15 and d["completed"] is False
    env.clock.advance(15)
    d = hb_confirm(sid)
    assert d["total_seconds"] == 30

    # 断点位置随心跳落库
    env.clock.advance(15)
    d = hb_confirm(sid)
    assert d["total_seconds"] == 45

    # 结束学习段；断点续学返回上次位置
    env.clock.advance(15)
    stopd = env.stop(sid, page=7)
    assert stopd["gained_seconds"] == 15 and stopd["total_seconds"] == 60
    data2 = env.start(cid)
    assert data2["total_seconds"] == 60 and data2["last_position"]["page"] == 7, "断点未保存/未恢复"
    sid2 = data2["session_id"]

    # 达标判定：60 分钟课程 → 快进到剩余秒数耗尽
    env.clock.advance(15)
    d = hb_confirm(sid2)
    assert d["total_seconds"] == 75
    for _ in range(140):  # 140*15=2100s + 75 = 2175，仍差 1425
        env.clock.advance(15)
        d = hb_confirm(sid2)
    assert d["completed"] is False and d["total_seconds"] == 2175
    for _ in range(95):  # 95*15=1425 → 3600 达标
        env.clock.advance(15)
        d = hb_confirm(sid2)
    assert d["completed"] is True and d["total_seconds"] == 3600, "达标判定失败: %s" % d

    # 达标后继续学习仍累计（超额如实记录），状态不回退
    env.clock.advance(15)
    d = hb_confirm(sid2)
    assert d["total_seconds"] == 3615 and d["completed"] is True

    # 学习明细只增不改：段计数与心跳数可对账
    env.stop(sid2)
    import sqlite3 as _s
    conn = _s.connect(str(env.settings.db_path))
    rows = conn.execute("SELECT gained_seconds, heartbeat_count, end_reason FROM study_sessions ORDER BY id").fetchall()
    rec = conn.execute("SELECT total_seconds, status, completed_at FROM study_records").fetchone()
    conn.close()
    assert sum(r[0] for r in rows) == rec[0] == 3615, "明细合计与累计口径不一致"
    assert rec[1] == "completed" and rec[2], "达标状态/达标时间未落库"
    assert all(r[2] in ("normal",) for r in rows), "结束原因异常: %s" % rows


@with_env
def test_max_gap_and_clock_rollback(env):
    env.standard_learner(courses=["通识/边界测试.pdf"])
    sid = env.start(env.course_id("通识/边界测试.pdf"))["session_id"]
    # 断线超长（61s > max_gap 60）→ 本段不计
    env.clock.advance(61)
    d = env.hb(sid)
    assert d["gained_seconds"] == 0 and d["total_seconds"] == 0
    # 恢复正常后继续计
    env.clock.advance(15)
    d = env.hb(sid)
    assert d["gained_seconds"] == 15
    # 时钟回拨/重放 → 不计，且回拨后时刻成为新的计时基准（保守不虚增）
    env.clock.advance(-30)
    d = env.hb(sid)
    assert d["gained_seconds"] == 0
    env.clock.advance(45)  # 自新基准起 45s
    d = env.hb(sid)
    assert d["gained_seconds"] == 45 and d["total_seconds"] == 60
    # 切屏（visible=false）→ 不计
    env.clock.advance(15)
    d = env.hb(sid, visible=False)
    assert d["gained_seconds"] == 0
    env.clock.advance(15)
    d = env.hb(sid, visible=True)
    assert d["gained_seconds"] == 15


@with_env
def test_challenge_state_machine(env):
    """验收标准 3：防挂机弹窗生效，不响应则计时暂停。"""
    study = {"heartbeat_interval_s": 15, "max_gap_s": 60,
             "challenge_min_s": 5, "challenge_max_s": 10, "challenge_timeout_s": 30,
             "daily_cap_hours": 8}
    e = Env(study=study)
    try:
        e.login_ok(ADMIN_NO, ADMIN_PW)
        e.import_roster([[LEARNER["employee_no"], LEARNER["name"], LEARNER["department"], ""]])
        e.put_course("通识/防挂机.pdf")
        e.scan()
        cid = e.course_id("通识/防挂机.pdf")
        e.client.post("/api/v1/auth/logout")
        e.login_ok("A001", "A001")
        sid = e.start(cid)["session_id"]

        # 1) 随机窗口内（5~10s）触发：advance 10s 保证触发
        e.clock.advance(10)
        d = e.hb(sid)
        assert d["gained_seconds"] == 10
        ch = d["challenge"]
        assert ch and ch["token"] and 0 < ch["timeout_s"] <= 30, "弹窗未触发: %s" % d
        token = ch["token"]

        # 2) 错误令牌 → 4202
        j = e.confirm(sid, "bad-token")
        assert j["code"] == 4202, j

        # 3) 正确确认 → 继续计时，无弹窗
        e.clock.advance(2)
        j = e.confirm(sid, token)
        assert j["code"] == 0 and j["data"]["resumed"] is False and j["data"]["confirm_ok"] == 1
        e.clock.advance(3)
        d = e.hb(sid)
        assert d["challenge"] is None and d["gained_seconds"] == 3 + 2  # 确认间隔 2s + 3s 连续计时

        # 4) 不响应弹窗 → 超时暂停：计时计到超时时刻为止
        e.clock.advance(10)  # 触发第二次弹窗
        d = e.hb(sid)
        assert d["challenge"] and d["challenge"]["token"] != token, "第二次弹窗未触发"
        token2 = d["challenge"]["token"]
        e.clock.advance(20)  # 20s < 30s 超时
        d = e.hb(sid)
        assert d["paused"] is False and d["challenge"]["token"] == token2
        e.clock.advance(11)  # 累计 31s > 30s 超时
        d = e.hb(sid)
        # 超时落在本心跳间隔中间：只计到超时时刻（本间隔 11s 中前 10s 有效）
        assert d["paused"] is True and d["gained_seconds"] == 10, "超时后未暂停: %s" % d
        assert d["challenge"]["timeout_s"] == 0
        # 暂停期间心跳一律 0（验收标准 3）
        e.clock.advance(15)
        d = e.hb(sid)
        assert d["gained_seconds"] == 0 and d["paused"] is True
        # 超时次数已入账（台账可见）
        e.clock.advance(15)
        e.hb(sid)

        # 5) 超时后的重新确认 → 恢复计时
        e.clock.advance(5)
        j = e.confirm(sid, token2)
        assert j["code"] == 0 and j["data"]["resumed"] is True, "重新确认未恢复: %s" % j
        e.clock.advance(15)
        d = e.hb(sid)
        assert d["paused"] is False and d["gained_seconds"] == 15, "恢复后计时未继续: %s" % d

        # 6) 令牌一次性：旧令牌再用 → 4202
        j = e.confirm(sid, token2)
        assert j["code"] == 4202

        # 7) confirm_miss / confirm_ok 落库
        conn = sqlite3.connect(str(e.settings.db_path))
        row = conn.execute("SELECT confirm_ok, confirm_miss FROM study_sessions WHERE id=?", (sid,)).fetchone()
        rec = conn.execute("SELECT challenge_miss FROM study_records").fetchone()
        conn.close()
        assert row[0] == 2 and row[1] == 1, "确认/超时计数不符: %s" % (row,)
        assert rec[0] == 1
    finally:
        e.close()


@with_env
def test_daily_cap(env):
    study = {"heartbeat_interval_s": 15, "max_gap_s": 60, "challenge_min_s": 10 ** 6,
             "challenge_max_s": 10 ** 6 + 1, "challenge_timeout_s": 120, "daily_cap_hours": 1}
    e = Env(study=study)
    try:
        e.login_ok(ADMIN_NO, ADMIN_PW)
        e.import_roster([[LEARNER["employee_no"], LEARNER["name"], LEARNER["department"], ""]])
        e.put_course("通识/上限测试.pdf")
        e.scan()
        cid = e.course_id("通识/上限测试.pdf")
        e.client.post("/api/v1/auth/logout")
        e.login_ok("A001", "A001")
        sid = e.start(cid)["session_id"]
        total = 0
        for i in range(300):  # 300*15=4500s > 3600s 上限
            e.clock.advance(15)
            d = e.hb(sid)
            total += d["gained_seconds"]
        assert total == 3600, "单日上限未生效: %d" % total
        assert d["total_seconds"] == 3600, "累计值超过单日上限: %s" % d
    finally:
        e.close()


@with_env
def test_sweep_and_double_tab_guard(env):
    study = {"heartbeat_interval_s": 15, "max_gap_s": 60, "challenge_min_s": 10 ** 6,
             "challenge_max_s": 10 ** 6 + 1, "challenge_timeout_s": 120, "daily_cap_hours": 8}
    e = Env(study=study)
    try:
        e.login_ok(ADMIN_NO, ADMIN_PW)
        e.import_roster([[LEARNER["employee_no"], LEARNER["name"], LEARNER["department"], ""]])
        e.put_course("通识/清理测试.pdf")
        e.scan()
        cid = e.course_id("通识/清理测试.pdf")
        e.client.post("/api/v1/auth/logout")
        e.login_ok("A001", "A001")
        # 双开防护：同一课程第二个学习段开启时，第一段自动封口
        sid1 = e.start(cid)["session_id"]
        e.clock.advance(15)
        e.hb(sid1)
        sid2 = e.start(cid)["session_id"]
        assert sid2 != sid1
        r = e.client.post("/api/v1/study/heartbeat", json={"session_id": sid1, "visible": True}).json()
        assert r["code"] == 4201, "旧学习段未封口: %s" % r
        # 心跳超时清理：超过 max_gap 无心跳 → timeout 封口
        e.clock.advance(15)
        e.hb(sid2)
        e.clock.advance(120)  # > max_gap 60
        closed = e.app.state.engine.sweep()
        assert closed == 1, "超时清理未生效: %d" % closed
        r = e.client.post("/api/v1/study/heartbeat", json={"session_id": sid2, "visible": True}).json()
        assert r["code"] == 4201
        conn = sqlite3.connect(str(e.settings.db_path))
        rows = conn.execute("SELECT end_reason FROM study_sessions ORDER BY id").fetchall()
        conn.close()
        assert [r[0] for r in rows] == ["timeout", "timeout"], rows
        # 重启封口：遗留开放段 → offline
        sid3 = e.start(cid)["session_id"]
        n = e.app.state.engine.close_stale_at_startup()
        assert n == 1
        conn = sqlite3.connect(str(e.settings.db_path))
        row = conn.execute("SELECT end_reason FROM study_sessions WHERE id=?", (sid3,)).fetchone()
        conn.close()
        assert row[0] == "offline"
    finally:
        e.close()


# ====================== 文件下载 ======================
@with_env
def test_file_download_range(env):
    pdf = make_pdf(pages=3)
    env.standard_learner(courses={"高处作业/下载测试.pdf": pdf})
    cid = env.course_id("高处作业/下载测试.pdf")
    # 全量下载
    r = env.client.get("/api/v1/courses/%d/file" % cid)
    assert r.status_code == 200 and r.content == pdf
    assert r.headers["accept-ranges"] == "bytes"
    assert "filename*=UTF-8''" in r.headers.get("content-disposition", "")
    # Range 分段（手机阅读器分页加载）
    r = env.client.get("/api/v1/courses/%d/file" % cid, headers={"Range": "bytes=0-99"})
    assert r.status_code == 206 and len(r.content) == 100
    assert r.headers["content-range"] == "bytes 0-99/%d" % len(pdf)
    r = env.client.get("/api/v1/courses/%d/file" % cid, headers={"Range": "bytes=100-199"})
    assert r.status_code == 206 and r.content == pdf[100:200]
    # 非法/不存在资源
    r = env.client.get("/api/v1/courses/9999/file")
    assert r.status_code == 404 and r.json()["code"] == 4041
    # 路径穿越安全：接口仅按 id 反查，不接受路径参数（此处验证不存在的 id 不泄露文件系统）
    r = env.client.get("/api/v1/courses/%d/file" % (cid + 100))
    assert r.status_code == 404


# ====================== 导出与总览 ======================
def _seed_export_env():
    study = {"heartbeat_interval_s": 15, "max_gap_s": 60, "challenge_min_s": 10 ** 6,
             "challenge_max_s": 10 ** 6 + 1, "challenge_timeout_s": 120, "daily_cap_hours": 8}
    e = Env(study=study)
    e.login_ok(ADMIN_NO, ADMIN_PW)
    e.import_roster([
        ["A001", "王建国", "一项目部", ""],
        ["A002", "张铁柱", "二项目部", ""],
        ["A003", "刘线务", "一项目部", ""],
    ])
    e.put_course("高处作业/登高作业安全要点.pdf")
    e.put_course("有限空间/人孔井作业安全规范.pdf")
    e.scan()
    e.client.post("/api/v1/auth/logout")
    # A001 学两门课（每门：3 次心跳 45s + 结束段 15s = 60s）
    e.login_ok("A001", "A001")
    for rel in ("高处作业/登高作业安全要点.pdf", "有限空间/人孔井作业安全规范.pdf"):
        sid = e.start(e.course_id(rel))["session_id"]
        for _ in range(3):
            e.clock.advance(15)
            e.hb(sid)
        e.clock.advance(15)   # 结束前最后一段：距上次心跳 15s，计入
        e.stop(sid)
    # A002 学一门（1 次心跳 15s + 结束段 15s = 30s）
    e.client.post("/api/v1/auth/logout")
    e.login_ok("A002", "A002")
    sid = e.start(e.course_id("高处作业/登高作业安全要点.pdf"))["session_id"]
    for _ in range(1):
        e.clock.advance(15)
        e.hb(sid)
    e.clock.advance(15)       # 结束前最后一段
    e.stop(sid)
    e.client.post("/api/v1/auth/logout")
    e.login_ok(ADMIN_NO, ADMIN_PW)
    return e


@with_env
def test_export_personal(env):
    e = _seed_export_env()
    try:
        # xlsx：三表结构
        r = e.client.get("/api/v1/admin/export/personal")
        assert r.status_code == 200 and "spreadsheetml" in r.headers["content-type"]
        wb = load_workbook(io.BytesIO(r.content))
        assert wb.sheetnames == ["个人学时台账", "学习段明细", "口径说明"], wb.sheetnames
        ws1 = wb["个人学时台账"]
        assert ws1.max_row == 4  # 表头 + A001×2 + A002×1（A003 未学习不出现在课程台账）
        hdr = [c.value for c in ws1[1]]
        assert hdr[:6] == ["工号", "姓名", "部门/班组", "课程", "课程分类", "累计学时(分钟)"]
        row_a1 = [c.value for c in ws1[2]]
        assert row_a1[0] == "A001" and row_a1[5] == "1.0"  # 60 秒 = 1.0 分钟
        ws2 = wb["学习段明细"]
        assert ws2.max_row == 4  # 表头 + 3 个学习段
        assert [c.value for c in ws2[1]][7] == "有效学时(分钟)"
        ws3 = wb["口径说明"]
        texts = " ".join(str(c.value) for row in ws3.iter_rows() for c in row if c.value)
        assert "服务端权威计时" in texts and "单日计入学时上限 8 小时" in texts, "口径说明附页缺失关键参数"
        # csv：utf-8-sig + 明细行
        r = e.client.get("/api/v1/admin/export/personal?format=csv")
        text = r.content.decode("utf-8-sig")
        lines = [l for l in text.splitlines() if l.strip()]
        assert lines[0].startswith("工号,姓名") and len(lines) == 4
        assert "A001" in lines[1]
        # 部门过滤
        r = e.client.get("/api/v1/admin/export/personal?department=%E4%B8%80%E9%A1%B9%E7%9B%AE%E9%83%A8")
        wb2 = load_workbook(io.BytesIO(r.content))
        ws = wb2["个人学时台账"]
        depts = {ws.cell(row=i, column=3).value for i in range(2, ws.max_row + 1)}
        assert depts == {"一项目部"} and ws.max_row == 3, "部门过滤失效: %s" % depts
        # 日期过滤（学习段明细按日过滤）
        today = e.clock.now.astimezone().strftime("%Y-%m-%d")
        r = e.client.get("/api/v1/admin/export/personal?date_from=%s&date_to=%s" % (today, today))
        wb3 = load_workbook(io.BytesIO(r.content))
        assert wb3["学习段明细"].max_row == 4, "当日过滤应包含全部 3 个学习段"
        r = e.client.get("/api/v1/admin/export/personal?date_from=2099-01-01")
        wb4 = load_workbook(io.BytesIO(r.content))
        assert wb4["学习段明细"].max_row == 1, "未来日期过滤应为空"
        # 非法日期 → 4001
        r = e.client.get("/api/v1/admin/export/personal?date_from=bad")
        assert r.json()["code"] == 4001
        # 导出留痕
        conn = sqlite3.connect(str(e.settings.db_path))
        logs = conn.execute("SELECT export_type, scope, row_count FROM export_logs ORDER BY id").fetchall()
        conn.close()
        assert len(logs) >= 5 and all(l[0] == "personal" for l in logs), logs
        j = e.client.get("/api/v1/admin/export/logs").json()
        assert j["code"] == 0 and len(j["data"]) >= 5
    finally:
        e.close()


@with_env
def test_export_summary(env):
    e = _seed_export_env()
    try:
        r = e.client.get("/api/v1/admin/export/summary")
        assert r.status_code == 200
        wb = load_workbook(io.BytesIO(r.content))
        assert wb.sheetnames == ["部门汇总", "口径说明"]
        ws = wb["部门汇总"]
        rows = [[c.value for c in row] for row in ws.iter_rows()]
        assert rows[0][0] == "部门/班组"
        data = {r[0]: r for r in rows[1:-1]}
        assert rows[-1][0] == "合计"
        # 一项目部：A001 两门 120s + A003 0s = 2 分钟；二项目部：A002 30s = 0.5 分钟
        assert data["一项目部"][2] == "2.0", data["一项目部"]
        assert data["二项目部"][2] == "0.5", data["二项目部"]
        assert data["一项目部"][1] == 2 and data["二项目部"][1] == 1
        assert rows[-1][2] == "2.5"
        # csv
        r = e.client.get("/api/v1/admin/export/summary?format=csv")
        text = r.content.decode("utf-8-sig")
        assert "部门/班组" in text and "一项目部" in text and "合计" in text
        # 学员不可导出
        e.client.post("/api/v1/auth/logout")
        e.login_ok("A001", "A001")
        r = e.client.get("/api/v1/admin/export/summary")
        assert r.json()["code"] == 4501
    finally:
        e.close()


@with_env
def test_overview(env):
    e = _seed_export_env()
    try:
        j = e.client.get("/api/v1/admin/overview").json()
        assert j["code"] == 0
        d = j["data"]
        assert d["total_users"] == 4  # admin + 3 学员
        assert d["active_courses"] == 2
        assert d["total_minutes"] == 2  # 150 秒 ≈ 2 分钟（整除）
        users = {u["employee_no"]: u for u in d["users"]}
        assert users["A001"]["course_count"] == 2 and users["A001"]["total_minutes"] == 2
        assert users["A002"]["course_count"] == 1 and users["A002"]["total_minutes"] == 0
        assert users["A003"]["course_count"] == 0 and users["A003"]["last_study_at"] == "—"
        assert set(d["departments"]) == {"一项目部", "二项目部", "安全管理中心"}
        # 部门过滤
        j = e.client.get("/api/v1/admin/overview?department=%E4%B8%80%E9%A1%B9%E7%9B%AE%E9%83%A8").json()
        assert all(u["department"] == "一项目部" for u in j["data"]["users"])
    finally:
        e.close()


@with_env
def test_backup_service(env):
    from app.services.backup import run_backup_now
    env.login_ok(ADMIN_NO, ADMIN_PW)
    env.put_course("x.pdf")
    env.scan()
    dest = run_backup_now(env.settings)
    assert Path(dest).is_file() and Path(dest).stat().st_size > 0, "备份文件未生成"
    # 备份可用 sqlite 打开且包含表
    conn = sqlite3.connect(dest)
    tables = {r[0] for r in conn.execute("SELECT name FROM sqlite_master WHERE type='table'")}
    conn.close()
    assert "users" in tables
    # 保留期清理：伪造过期备份
    old = env.settings.backups_dir / "app-backup-20200101-000000.db"
    old.write_bytes(b"x")
    import os
    import time
    os.utime(str(old), (time.time() - 40 * 86400, time.time() - 40 * 86400))
    run_backup_now(env.settings)
    assert not old.exists(), "过期备份未清理"


def test_learner_page_no_external_deps():
    for page in ("learner.html", "admin.html"):
        html = (REPO_ROOT / "static" / page).read_text(encoding="utf-8")
        assert "http://" not in html.replace("http://127.0.0.1", ""), "%s 含外部依赖" % page
        assert "https://" not in html, "%s 含外部依赖" % page
        assert "<script src" not in html, "%s 引用外部脚本" % page


def main():
    tests = [v for k, v in sorted(globals().items())
             if k.startswith("test_") and callable(v)]
    failed = 0
    for t in tests:
        try:
            t()
            print("PASS %s" % t.__name__)
        except AssertionError as e:
            failed += 1
            print("FAIL %s: %s" % (t.__name__, e))
        except Exception as e:
            failed += 1
            print("ERROR %s: %r" % (t.__name__, e))
    print("total=%d failed=%d" % (len(tests), failed))
    sys.exit(1 if failed else 0)


if __name__ == "__main__":
    main()
