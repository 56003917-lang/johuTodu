# -*- coding: utf-8 -*-
"""学时台账导出（产品设计文档 §3.4 / 技术设计文档 §5、§13）。

- 个人学时明细：课程、累计时长、防挂机确认记录，逐条可追溯（xlsx 三表 / csv 明细）；
- 部门/班组汇总：按项目部、班组汇总学时完成情况；
- 口径说明附页：计时参数表写进台账（口径透明是最好的抗质疑手段）；
- 每次导出在 export_logs 落一条记录（谁导的、导了什么、多少行）。
"""
import csv
import io
import logging
from datetime import datetime
from typing import Optional

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from sqlalchemy import func

from app.config import Settings
from app.db import Database
from app.models import Course, ExportLog, StudyRecord, StudySession, User
from app.util import date_range_utc, fmt_local, iso, today_window_utc, utcnow

log = logging.getLogger("app.exporter")

_HEADER_FILL = PatternFill("solid", fgColor="C9D6F3")
_TIP_FONT = Font(color="5B6B80", size=10)

REASON_CN = {"normal": "正常结束", "timeout": "心跳超时", "logout": "退出登录", "offline": "服务重启", "": ""}


def _fmt_minutes(seconds) -> str:
    """秒 → 学时（分钟），保留 1 位小数：台账以分钟为口径。"""
    return "%.1f" % (int(seconds or 0) / 60.0)


def _scope_text(department: Optional[str], date_from: Optional[str], date_to: Optional[str]) -> str:
    return "department=%s;from=%s;to=%s" % (department or "全部", date_from or "-", date_to or "-")


def _parse_filters(department, date_from, date_to):
    start, end = date_range_utc(date_from, date_to)
    if (date_from and not start) or (date_to and not end):
        raise ValueError("日期参数格式应为 YYYY-MM-DD")
    return (department.strip() if department and department.strip() else None), start, end


# ======================================================================
# 个人学时台账
# ======================================================================
def _query_records(s, department, start, end):
    q = (s.query(StudyRecord, User, Course)
         .join(User, StudyRecord.user_id == User.id)
         .join(Course, StudyRecord.course_id == Course.id)
         .filter(User.status == 1))
    if department:
        q = q.filter(User.department == department)
    q = q.order_by(User.department, User.employee_no, Course.title)
    return q.all()


def _query_sessions(s, department, start, end):
    q = (s.query(StudySession, User, Course)
         .join(User, StudySession.user_id == User.id)
         .join(Course, StudySession.course_id == Course.id)
         .filter(User.status == 1))
    if department:
        q = q.filter(User.department == department)
    if start:
        q = q.filter(StudySession.started_at >= iso(start))
    if end:
        q = q.filter(StudySession.started_at < iso(end))
    q = q.order_by(User.department, User.employee_no, StudySession.started_at)
    return q.all()


def _sheet_records(ws, rows):
    headers = ["工号", "姓名", "部门/班组", "课程", "课程分类", "累计学时(分钟)",
               "达标学时(分钟)", "达标状态", "达标时间", "累计防挂机超时", "首次学习", "最近学习"]
    ws.append(headers)
    for rec, u, c in rows:
        ws.append([
            u.employee_no, u.name, u.department, c.title, c.category,
            _fmt_minutes(rec.total_seconds), "%.0f" % c.required_minutes,
            "已达标" if rec.status == "completed" else "学习中",
            fmt_local(rec.completed_at) or "—",
            rec.challenge_miss,
            fmt_local(rec.first_started_at), fmt_local(rec.last_study_at) or "—",
        ])
    return len(rows)


def _sheet_sessions(ws, rows):
    headers = ["工号", "姓名", "部门/班组", "课程", "学习日期", "开始时间", "结束时间",
               "有效学时(分钟)", "心跳次数", "防挂机确认成功", "防挂机超时", "结束原因"]
    ws.append(headers)
    for sess, u, c in rows:
        ws.append([
            u.employee_no, u.name, u.department, c.title,
            fmt_local(sess.started_at, "%Y-%m-%d"), fmt_local(sess.started_at), fmt_local(sess.ended_at) or "进行中",
            _fmt_minutes(sess.gained_seconds), sess.heartbeat_count,
            sess.confirm_ok, sess.confirm_miss, REASON_CN.get(sess.end_reason or "", sess.end_reason or ""),
        ])
    return len(rows)


def _sheet_caliber(ws, settings: Settings, department, date_from, date_to):
    """口径说明附页（§13 风险对策：计时口径写进导出台账）。"""
    ws.append(["学时计时口径说明（导出时间 %s）" % datetime.now().strftime("%Y-%m-%d %H:%M")])
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=2)
    ws["A1"].font = Font(bold=True)
    st = settings.study
    items = [
        ("计时方式", "服务端权威计时：以收到心跳的服务端时间差为唯一依据，不信任前端时钟"),
        ("心跳间隔", "%d 秒/次" % st.heartbeat_interval_s),
        ("断线容忍", "心跳间隔超过 %d 秒的部分不计学时" % st.max_gap_s),
        ("切屏口径", "页面切至后台/不可见期间不计学时"),
        ("防挂机弹窗", "随机 %d~%d 秒触发；%d 秒内未确认则暂停计时，确认后恢复"
         % (st.challenge_min_s, st.challenge_max_s, st.challenge_timeout_s)),
        ("单日上限", "每人单日计入学时上限 %.0f 小时" % st.daily_cap_hours),
        ("达标判定", "累计有效秒数 ≥ 课程达标学时×60，服务端判定"),
        ("数据留痕", "学习明细只增不改；本台账由系统自动生成，不可手工修改"),
        ("导出范围", "部门=%s；日期=%s ~ %s" % (department or "全部", date_from or "-", date_to or "-")),
    ]
    for k, v in items:
        ws.append([k, v])
    ws.column_dimensions["A"].width = 16
    ws.column_dimensions["B"].width = 80


def _style_sheet(ws, n_cols: int, widths=None):
    for c in ws[1]:
        c.font = Font(bold=True)
        c.fill = _HEADER_FILL
        c.alignment = Alignment(horizontal="center")
    ws.freeze_panes = "A2"
    if widths:
        for i, w in enumerate(widths, start=1):
            ws.column_dimensions[get_column_letter(i)].width = w
    else:
        for i in range(1, n_cols + 1):
            ws.column_dimensions[get_column_letter(i)].width = 14


def export_personal(db: Database, settings: Settings, operator_id: int,
                    department=None, date_from=None, date_to=None, fmt="xlsx"):
    """个人学时台账：xlsx=三表（台账/学习明细/口径说明）；csv=学习段明细。"""
    department, start, end = _parse_filters(department, date_from, date_to)
    with db.session() as s:
        rec_rows = _query_records(s, department, None, None)   # 台账口径为累计（不受日期过滤影响）
        sess_rows = _query_sessions(s, department, start, end)

        if fmt == "csv":
            buf = io.StringIO()
            writer = csv.writer(buf)
            writer.writerow(["工号", "姓名", "部门/班组", "课程", "学习日期", "开始时间", "结束时间",
                             "有效学时(分钟)", "心跳次数", "防挂机确认成功", "防挂机超时", "结束原因"])
            for sess, u, c in sess_rows:
                writer.writerow([u.employee_no, u.name, u.department, c.title,
                                 fmt_local(sess.started_at, "%Y-%m-%d"), fmt_local(sess.started_at),
                                 fmt_local(sess.ended_at) or "进行中",
                                 _fmt_minutes(sess.gained_seconds), sess.heartbeat_count,
                                 sess.confirm_ok, sess.confirm_miss,
                                 REASON_CN.get(sess.end_reason or "", "")])
            content = ("\ufeff" + buf.getvalue()).encode("utf-8")  # BOM：Excel 直接打开不乱码
            filename = "个人学时明细_%s.csv" % datetime.now().strftime("%Y%m%d-%H%M")
            row_count = len(sess_rows)
        else:
            wb = Workbook()
            ws1 = wb.active
            ws1.title = "个人学时台账"
            _sheet_records(ws1, rec_rows)
            _style_sheet(ws1, 12, [12, 10, 22, 24, 14, 14, 14, 10, 18, 14, 18, 18])
            ws2 = wb.create_sheet("学习段明细")
            _sheet_sessions(ws2, sess_rows)
            _style_sheet(ws2, 12, [12, 10, 22, 24, 12, 18, 18, 14, 10, 14, 12, 12])
            ws3 = wb.create_sheet("口径说明")
            _sheet_caliber(ws3, settings, department, date_from, date_to)
            out = io.BytesIO()
            wb.save(out)
            content = out.getvalue()
            filename = "个人学时台账_%s.xlsx" % datetime.now().strftime("%Y%m%d-%H%M")
            row_count = len(rec_rows) + len(sess_rows)

        s.add(ExportLog(operator_id=operator_id, export_type="personal",
                        scope=_scope_text(department, date_from, date_to),
                        row_count=row_count, created_at=iso(utcnow())))
        s.commit()
    log.info("导出个人台账: rows=%d scope=%s", row_count, _scope_text(department, date_from, date_to))
    return filename, content, row_count


# ======================================================================
# 部门/班组汇总
# ======================================================================
def export_summary(db: Database, settings: Settings, operator_id: int,
                   department=None, date_from=None, date_to=None, fmt="xlsx"):
    department, start, end = _parse_filters(department, date_from, date_to)
    with db.session() as s:
        users_q = s.query(User).filter(User.status == 1)
        if department:
            users_q = users_q.filter(User.department == department)
        users = users_q.all()
        user_ids = [u.id for u in users]
        dept_users = {}
        for u in users:
            dept_users.setdefault(u.department or "未分配", []).append(u)

        # 时长与超时（可按日期过滤的学习段口径）
        sess_q = s.query(StudySession)
        if user_ids:
            sess_q = sess_q.filter(StudySession.user_id.in_(user_ids))
        else:
            sess_q = sess_q.filter(StudySession.user_id == -1)
        if start:
            sess_q = sess_q.filter(StudySession.started_at >= iso(start))
        if end:
            sess_q = sess_q.filter(StudySession.started_at < iso(end))
        sess_rows = sess_q.all()
        sess_by_user = {}
        for r in sess_rows:
            sess_by_user.setdefault(r.user_id, []).append(r)

        # 达标人次（达标时间在范围内；未给范围则累计口径）
        comp_q = (s.query(StudyRecord).filter(StudyRecord.status == "completed"))
        if user_ids:
            comp_q = comp_q.filter(StudyRecord.user_id.in_(user_ids))
        else:
            comp_q = comp_q.filter(StudyRecord.user_id == -1)
        if start:
            comp_q = comp_q.filter(StudyRecord.completed_at >= iso(start))
        if end:
            comp_q = comp_q.filter(StudyRecord.completed_at < iso(end))
        completed_by_user = {}
        for r in comp_q.all():
            completed_by_user[r.user_id] = completed_by_user.get(r.user_id, 0) + 1

        header = ["部门/班组", "在册人数", "累计有效学时(分钟)", "人均学时(分钟)",
                  "达标人次", "防挂机超时合计", "参与学习人数"]
        table = []
        for dept in sorted(dept_users.keys()):
            ulist = dept_users[dept]
            ids = [u.id for u in ulist]   # user 主键；ulist 存的是 User 对象，不能直接当 dict key
            total_s = sum(r.gained_seconds for uid in ids for r in sess_by_user.get(uid, []))
            miss = sum(r.confirm_miss for uid in ids for r in sess_by_user.get(uid, []))
            done = sum(completed_by_user.get(uid, 0) for uid in ids)
            learners = sum(1 for uid in ids if sess_by_user.get(uid))
            table.append([dept, len(ulist), _fmt_minutes(total_s),
                          "%.1f" % (total_s / 60.0 / max(1, len(ulist))), done, miss, learners])
        total_row = ["合计", len(users),
                     _fmt_minutes(sum(r.gained_seconds for rs in sess_by_user.values() for r in rs)),
                     "%.1f" % (sum(r.gained_seconds for rs in sess_by_user.values() for r in rs) / 60.0 / max(1, len(users))),
                     sum(completed_by_user.values()),
                     sum(r.confirm_miss for rs in sess_by_user.values() for r in rs),
                     sum(1 for uid in user_ids if sess_by_user.get(uid))]

        if fmt == "csv":
            buf = io.StringIO()
            writer = csv.writer(buf)
            writer.writerow(header)
            writer.writerows(table)
            writer.writerow(total_row)
            content = ("\ufeff" + buf.getvalue()).encode("utf-8")
            filename = "部门学时汇总_%s.csv" % datetime.now().strftime("%Y%m%d-%H%M")
            row_count = len(table)
        else:
            wb = Workbook()
            ws1 = wb.active
            ws1.title = "部门汇总"
            ws1.append(header)
            for r in table:
                ws1.append(r)
            ws1.append(total_row)
            _style_sheet(ws1, len(header), [26, 12, 20, 16, 12, 18, 14])
            for c in ws1[ws1.max_row]:
                c.font = Font(bold=True)
            ws2 = wb.create_sheet("口径说明")
            _sheet_caliber(ws2, settings, department, date_from, date_to)
            out = io.BytesIO()
            wb.save(out)
            content = out.getvalue()
            filename = "部门学时汇总_%s.xlsx" % datetime.now().strftime("%Y%m%d-%H%M")
            row_count = len(table)

        s.add(ExportLog(operator_id=operator_id, export_type="summary",
                        scope=_scope_text(department, date_from, date_to),
                        row_count=row_count, created_at=iso(utcnow())))
        s.commit()
    log.info("导出部门汇总: rows=%d scope=%s", row_count, _scope_text(department, date_from, date_to))
    return filename, content, row_count


# ======================================================================
# 管理端总览（首页表格数据）
# ======================================================================
def overview(db: Database, department=None) -> dict:
    with db.session() as s:
        from sqlalchemy import case
        q = (s.query(
            User.id, User.employee_no, User.name, User.department, User.role,
            func.coalesce(func.sum(StudyRecord.total_seconds), 0),
            func.coalesce(func.count(StudyRecord.id), 0),
            func.coalesce(func.sum(case((StudyRecord.status == "completed", 1), else_=0)), 0),
            func.max(StudyRecord.last_study_at),
        ).outerjoin(StudyRecord, StudyRecord.user_id == User.id)
             .filter(User.status == 1))
        if department and department.strip():
            q = q.filter(User.department == department.strip())
        rows = q.group_by(User.id).order_by(User.department, User.employee_no).all()

        start, end = today_window_utc()
        today_q = (s.query(StudySession.user_id,
                           func.coalesce(func.sum(StudySession.gained_seconds), 0))
                   .filter(StudySession.started_at >= iso(start),
                           StudySession.started_at < iso(end))
                   .group_by(StudySession.user_id))
        today_map = {uid: int(v or 0) for uid, v in today_q.all()}

        active_courses = s.query(func.count(Course.id)).filter(Course.status == 1).scalar() or 0
        departments = [d[0] for d in s.query(User.department).filter(
            User.status == 1, User.department != "").distinct().order_by(User.department).all()]

        users = [{
            "employee_no": r[1], "name": r[2], "department": r[3], "role": r[4],
            "total_minutes": int(r[5] or 0) // 60,
            "course_count": int(r[6] or 0),
            "completed_count": int(r[7] or 0),
            "last_study_at": fmt_local(r[8]) if r[8] else "—",
            "today_minutes": today_map.get(r[0], 0) // 60,
        } for r in rows]
        return {
            "users": users,
            "departments": departments,
            "total_users": len(users),
            "total_minutes": sum(u["total_minutes"] for u in users),
            "today_minutes": sum(u["today_minutes"] for u in users),
            "active_courses": int(active_courses),
        }
