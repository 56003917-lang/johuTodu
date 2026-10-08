# -*- coding: utf-8 -*-
"""名册服务：Excel 批量导入员工名册 + 模板下载（产品设计文档 §3.1）。

导入即生成账号：账号=工号，初始密码默认为工号（策略可配置）；
重复工号即更新而非报错（软删除：status 恢复为在册，历史学时保留）。
"""
import io
import logging
from typing import Dict

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill

from app.config import Settings
from app.db import Database
from app.security import hash_password

log = logging.getLogger("app.roster")

TEMPLATE_HEADERS = ["工号", "姓名", "部门/班组", "角色"]
ROLE_MAP = {"管理员": "admin", "admin": "admin", "学员": "learner", "learner": "learner", "": "learner"}


def _cell_str(v) -> str:
    """Excel 单元格值 → 干净字符串（处理数字/浮点/空格）。"""
    if v is None:
        return ""
    if isinstance(v, float):
        return str(int(v)) if v.is_integer() else str(v)
    if isinstance(v, int):
        return str(v)
    return str(v).strip()


def generate_template() -> bytes:
    """名册 Excel 模板（表头 + 示例行 + 填写说明）。"""
    wb = Workbook()
    ws = wb.active
    ws.title = "员工名册"
    header_fill = PatternFill("solid", fgColor="C9D6F3")
    ws.append(TEMPLATE_HEADERS)
    for c in ws[1]:
        c.font = Font(bold=True)
        c.fill = header_fill
        c.alignment = Alignment(horizontal="center")
    ws.append(["A001", "张三", "一项目部·高处作业班组", "学员"])
    ws.append(["B002", "李四", "安全管理中心", "管理员"])
    ws.append(["C003", "王五", "二项目部·线路班组", ""])  # 角色留空默认学员
    widths = [14, 12, 30, 10]
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[chr(64 + i)].width = w
    ws.merge_cells("A6:D6")
    tip = ws.cell(row=6, column=1, value="说明：角色列填“管理员/学员”，留空默认学员；重复工号再次导入即更新信息，不改密码。")
    tip.font = Font(color="5B6B80", size=10)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def import_roster(db: Database, settings: Settings, data: bytes) -> Dict[str, int]:
    """导入名册（新增+更新），返回统计。"""
    try:
        wb = load_workbook(io.BytesIO(data), data_only=True, read_only=True)
    except Exception:
        raise ValueError("不是有效的 xlsx 文件")
    ws = wb.active
    rows = list(ws.iter_rows(values_only=True))
    wb.close()
    # 容错：首行是表头则跳过
    if rows and _cell_str(rows[0][0]) == "工号":
        rows = rows[1:]

    from app.models import User
    from app.util import iso, utcnow
    now = iso(utcnow())
    created = updated = skipped = 0
    seen = set()
    with db.session() as s:
        for idx, row in enumerate(rows, start=2):
            if not row or all(v is None or _cell_str(v) == "" for v in row):
                continue
            employee_no = _cell_str(row[0])
            name = _cell_str(row[1]) if len(row) > 1 else ""
            department = _cell_str(row[2]) if len(row) > 2 else ""
            role_text = _cell_str(row[3]) if len(row) > 3 else ""
            if not employee_no or not name:
                skipped += 1
                log.warning("名册第 %d 行缺少工号或姓名，已跳过", idx)
                continue
            if employee_no in seen:
                skipped += 1
                continue
            seen.add(employee_no)
            role = ROLE_MAP.get(role_text.lower(), "learner")
            user = s.query(User).filter(User.employee_no == employee_no).one_or_none()
            if user:
                # 更新而非报错（§4.2）；不重置密码，离职复活
                user.name = name
                user.department = department
                user.role = role
                user.status = 1
                user.updated_at = now
                updated += 1
            else:
                s.add(User(
                    employee_no=employee_no, name=name, department=department, role=role,
                    password_hash=hash_password(settings.initial_password_for(employee_no)),
                    must_change_pwd=1 if settings.force_change_password else 0,
                    status=1, created_at=now, updated_at=now,
                ))
                created += 1
        s.commit()
    total = created + updated + skipped
    log.info("名册导入: 共 %d 行，新增 %d，更新 %d，跳过 %d", total, created, updated, skipped)
    return {"total": total, "created": created, "updated": updated, "skipped": skipped}
